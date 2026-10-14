import csv
import io
import re
import zipfile
from bisect import bisect_left, bisect_right
from pathlib import Path
from .base import ExtractedSection, SourceError, SourceSnapshot
from ..ingest_limits import (
    MAX_ARCHIVE_ENTRIES,
    MAX_IMAGE_PIXELS,
    MAX_OFFICE_EXPANDED_BYTES,
    MAX_PDF_PAGES,
    MAX_SMALL_SOURCE_BYTES,
    MAX_STRUCTURED_DEPTH,
    MAX_STRUCTURED_VALUES,
    MAX_TEXT_CELLS,
    MAX_DOCX_BLOCKS,
    MAX_OCR_PAGES_PER_BATCH,
    MAX_TEACHING_WINDOWS,
)

SUPPORTED = {
    '.pdf', '.md', '.markdown', '.txt', '.docx', '.rst', '.csv', '.tsv',
    '.xlsx', '.xlsm', '.json', '.yml', '.yaml', '.log', '.png', '.jpg', '.jpeg',
}
MAX_TEXT_CHUNK_CHARS = 4000
TEACHING_WINDOW_TOKENS = 160
TEACHING_CONTEXT_OVERLAP_TOKENS = 20


def split_teaching_windows(text: str) -> list[dict[str, str | int]]:
    """Return exact source slices in 160-token windows with 20-token overlap.

    Window ends prefer paragraph boundaries, then sentence boundaries, then
    whitespace, matching the source project's SentenceSplitter policy. Token
    boundaries are converted back to safe UTF-8 character offsets so citations
    remain exact excerpts, including for CJK text and emoji.

    ``core_text`` excludes the preceding overlap and is the unique assessed
    concept. ``text`` includes the overlap as context for the tutor/citation.
    """
    import tiktoken

    if not text or not text.strip():
        return []
    encoder = tiktoken.get_encoding("cl100k_base")
    tokens = encoder.encode(text, disallowed_special=())
    if not tokens:
        return []
    raw = text.encode("utf-8")
    token_bytes = [0]
    for token in tokens:
        token_bytes.append(token_bytes[-1] + len(encoder.decode_single_token_bytes(token)))
    if token_bytes[-1] != len(raw):
        raise ValueError("Tokenizer byte offsets did not resolve to the original text")

    char_bytes = [0]
    for char in text:
        char_bytes.append(char_bytes[-1] + len(char.encode("utf-8")))

    def char_at(token_index: int) -> int:
        """Nearest valid character boundary for a possibly partial UTF-8 token."""
        byte_offset = token_bytes[token_index]
        right = bisect_left(char_bytes, byte_offset)
        if right >= len(char_bytes):
            return len(text)
        if char_bytes[right] == byte_offset or right == 0:
            return right
        left = right - 1
        return left if byte_offset - char_bytes[left] <= char_bytes[right] - byte_offset else right

    token_chars = [char_at(i) for i in range(len(token_bytes))]
    separator_patterns = (
        re.compile(r"\n[ \t]*\n+"),
        re.compile(r"(?<=[.!?。！？])(?:[ \t]+|\n+|$)"),
        re.compile(r"\s+"),
    )
    results: list[dict[str, str | int]] = []
    start_token = 0
    previous_end_token = 0
    while start_token < len(tokens):
        start_char = token_chars[start_token]
        target_token = min(start_token + TEACHING_WINDOW_TOKENS, len(tokens))
        target_char = token_chars[target_token]
        end_token = target_token

        # Stay within the token budget while preferring the strongest available
        # boundary. Ignore a boundary that would create an empty/tiny window.
        for pattern in separator_patterns:
            matches = [m for m in pattern.finditer(text, start_char, target_char) if m.end() > start_char]
            if matches:
                chosen_char = matches[-1].end()
                candidate = bisect_right(token_chars, chosen_char) - 1
                if candidate > start_token:
                    end_token = min(candidate, target_token)
                    break

        while end_token > start_token + 1:
            end_char = token_chars[end_token]
            if len(encoder.encode(text[start_char:end_char], disallowed_special=())) <= TEACHING_WINDOW_TOKENS:
                break
            end_token -= 1
        end_char = token_chars[end_token]
        if end_char <= start_char:
            end_token = min(start_token + 1, len(tokens))
            end_char = token_chars[end_token]

        window_start = start_char
        window_end = end_char
        while window_start < window_end and text[window_start].isspace():
            window_start += 1
        while window_end > window_start and text[window_end - 1].isspace():
            window_end -= 1

        core_start = token_chars[previous_end_token] if results else window_start
        core_start = max(window_start, core_start)
        while core_start < window_end and text[core_start].isspace():
            core_start += 1
        window_text = text[window_start:window_end]
        core_text = text[core_start:window_end]
        if window_text and core_text:
            results.append({
                "text": window_text,
                "core_text": core_text,
                "start": window_start,
                "end": window_end,
                "core_start": core_start,
                "core_end": window_end,
                "token_start": start_token,
                "token_end": end_token,
                "core_token_start": previous_end_token if results else start_token,
            })
        if end_token == len(tokens):
            break
        previous_end_token = end_token
        start_token = max(start_token + 1, end_token - TEACHING_CONTEXT_OVERLAP_TOKENS)
    return results


def coherent_sections(section: ExtractedSection) -> list[ExtractedSection]:
    """Bound retrieval units within a source section, preferring real boundaries.

    Never cross a heading or PDF page. Offsets address the original extracted
    section, and every resulting excerpt is an exact contiguous source slice.
    """
    text = section.text
    if len(text) <= MAX_TEXT_CHUNK_CHARS:
        return [section]
    result, start = [], 0
    while start < len(text):
        end = min(start + MAX_TEXT_CHUNK_CHARS, len(text))
        if end < len(text):
            window = text[start:end]
            boundaries = list(re.finditer(r'\n\s*\n', window))
            if not boundaries:
                boundaries = list(re.finditer(r'[.!?。！？](?:\s+|(?=[^\x00-\x7f]))', window))
            if not boundaries:
                boundaries = list(re.finditer(r'\s+', window))
            if boundaries:
                end = start + boundaries[-1].end()
        segment = text[start:end]
        left = start + len(segment) - len(segment.lstrip())
        right = end - len(segment) + len(segment.rstrip())
        if left < right:
            base_start = int(section.locator.get('section_char_start', 0))
            locator = {**section.locator, 'section_char_start': base_start + left, 'section_char_end': base_start + right}
            if section.locator_kind == 'text':
                locator['line_start'] = section.locator['line_start'] + text[:left].count('\n')
                locator['line_end'] = section.locator['line_start'] + text[:right].count('\n')
            result.append(ExtractedSection(f'{section.title} — {len(result)+1}', text[left:right], section.locator_kind, locator))
        start = end
    return result


def _normalise_pdf_text(text: str) -> str:
    lines = [re.sub(r"[ \t]{2,}", " ", line).rstrip() for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines))


def _all_caps_heading(line: str) -> bool:
    value = line.strip()
    return bool(re.fullmatch(r"[A-Z][A-Z0-9 &/()'-]{2,59}", value))


def _prose_sections(text: str, filename: str, locator_kind: str, base_locator: dict, *, markdown=False, rst=False, all_caps=False) -> list[ExtractedSection]:
    """Split a prose source at its own headings and retain exact text offsets."""
    lines = text.splitlines(keepends=True)
    line_starts = [0]
    for line in lines:
        line_starts.append(line_starts[-1] + len(line))
    if len(lines) == 0:
        return []

    caps_count = sum(_all_caps_heading(line) for line in lines) if all_caps else 0
    nonempty = sum(bool(line.strip()) for line in lines)
    promote_caps = bool(all_caps and not (caps_count > 12 and caps_count * 4 > nonempty))
    heading_stack: list[tuple[int, str]] = []
    starts: list[tuple[int, int, str]] = []
    index = 0
    while index < len(lines):
        value = lines[index].strip()
        level = 0
        title = ""
        consumed = 1
        markdown_heading = re.match(r"^(#{1,6})\s+(\S.*)$", value) if markdown else None
        if markdown_heading:
            level, title = len(markdown_heading.group(1)), markdown_heading.group(2).strip()
        elif rst and index + 1 < len(lines):
            underline = re.fullmatch(r"([=\-`:'\"~^_*+#])\1{2,}\s*", lines[index + 1].strip())
            if value and underline and len(lines[index + 1].strip()) >= len(value):
                level, title, consumed = 2, value, 2
                if index > 0 and re.fullmatch(r"([=\-`:'\"~^_*+#])\1{2,}\s*", lines[index - 1].strip()):
                    starts.append((index - 1, 1, ""))
        elif promote_caps and _all_caps_heading(value):
            level, title = 2, value
        if level:
            starts.append((index, level, title))
            index += consumed
            continue
        index += 1

    result: list[ExtractedSection] = []
    body_start = 0
    heading_path: list[tuple[int, str]] = []
    current_title = Path(filename).stem[:160] or "Source"

    def append_body(end_line: int) -> None:
        nonlocal body_start
        raw_start = line_starts[body_start]
        raw_end = line_starts[end_line]
        raw = text[raw_start:raw_end]
        stripped = raw.strip()
        if not stripped:
            return
        leading = len(raw) - len(raw.lstrip())
        trailing = len(raw) - len(raw.rstrip())
        start_char = raw_start + leading
        end_char = raw_end - trailing
        body = text[start_char:end_char]
        line_start = 1 + text[:start_char].count("\n")
        line_end = 1 + text[:end_char].count("\n")
        path = [part for _, part in heading_path]
        locator = {
            **base_locator,
            "filename": filename,
            "heading": " / ".join(path) if path else current_title,
            "heading_path": path,
            "line_start": line_start,
            "line_end": max(line_start, line_end),
            "section_char_start": start_char,
            "section_char_end": end_char,
        }
        result.append(ExtractedSection(locator["heading"], body, locator_kind, locator))

    for heading_line, level, title in starts:
        if not title:
            continue
        append_body(heading_line)
        while heading_path and heading_path[-1][0] >= level:
            heading_path.pop()
        heading_path.append((level, title))
        current_title = title
        body_start = heading_line + (2 if rst else 1)
    append_body(len(lines))
    if result:
        return result

    # No headings: keep one coherent source section, using the first non-empty
    # PDF line only as a deterministic display label.
    body = text.strip()
    if not body:
        return []
    first_line = next((line.strip() for line in text.splitlines() if line.strip()), current_title)
    start_char = len(text) - len(text.lstrip())
    end_char = len(text.rstrip())
    line_start = 1 + text[:start_char].count("\n")
    line_end = 1 + text[:end_char].count("\n")
    locator = {
        **base_locator,
        "filename": filename,
        "heading": first_line[:160] if locator_kind == "pdf" else current_title,
        "heading_path": [],
        "line_start": line_start,
        "line_end": max(line_start, line_end),
        "section_char_start": start_char,
        "section_char_end": end_char,
    }
    return [ExtractedSection(locator["heading"], text[start_char:end_char], locator_kind, locator)]


def pdf_manifest_info(data: bytes) -> dict:
    """Read bounded PDF page count and bookmark paths without extracting page text."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise SourceError("encrypted_pdf", "Upload an unlocked PDF.")
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            raise SourceError("pdf_page_limit", f"Limit PDFs to {MAX_PDF_PAGES} pages.")
        bookmarks: list[tuple[int, list[str]]] = []

        def walk(items, parents: tuple[str, ...] = ()):
            current = parents
            for item in items or []:
                if isinstance(item, list):
                    walk(item, current)
                    continue
                title = str(getattr(item, "title", "")).strip()
                path = (*parents, title) if title else parents
                if title:
                    try:
                        page = reader.get_destination_page_number(item) + 1
                        if 1 <= page <= page_count:
                            bookmarks.append((page, list(path)))
                    except Exception:
                        pass
                children = getattr(item, "children", None)
                if children:
                    walk(children, path)
                current = path

        try:
            walk(reader.outline)
        except Exception:
            # Malformed/unsupported bookmarks do not prevent page extraction.
            bookmarks = []
        bookmarks.sort(key=lambda pair: pair[0])
        paths: list[list[str]] = []
        active: list[str] = []
        index = 0
        for page_number in range(1, page_count + 1):
            while index < len(bookmarks) and bookmarks[index][0] <= page_number:
                active = bookmarks[index][1]
                index += 1
            paths.append(list(active))
        return {"page_count": page_count, "chapter_paths": paths}
    except SourceError:
        raise
    except Exception as exc:
        raise SourceError("unreadable_pdf", "The PDF could not be read safely.") from exc


def _blank_page_image(image_bytes: bytes) -> bool:
    """Conservative blank-page heuristic; uncertain pages go to learner review."""
    from PIL import Image

    image = Image.open(io.BytesIO(image_bytes)).convert("L")
    image.thumbnail((128, 128))
    histogram = image.histogram()
    pixels = max(1, image.width * image.height)
    non_background = sum(histogram[:242])
    return non_background / pixels < 0.0005


def _page_text_warning(text: str, method: str) -> str | None:
    value = text.strip()
    if not value:
        return "empty_extraction"
    controls = sum(ord(char) < 32 and char not in "\n\r\t" for char in value)
    if "\ufffd" in value or controls / max(len(value), 1) > 0.01:
        return "suspicious_characters"
    if method == "vision":
        visible = sum(char.isalnum() or char.isspace() or char in ".,;:!?'-()[]" for char in value)
        if len(value) < 16 or visible / max(len(value), 1) < 0.55:
            return "low_ocr_quality"
    return None


def extract_pdf_page_batch(
    data: bytes,
    filename: str,
    start_page: int,
    end_page: int,
    chapter_paths: list[list[str]],
    vision=None,
) -> tuple[list[ExtractedSection], list[dict]]:
    """Extract one bounded, globally page-numbered batch and report each page."""
    from pypdf import PdfReader

    if end_page - start_page + 1 > MAX_OCR_PAGES_PER_BATCH:
        raise SourceError("pdf_batch_limit", f"Process at most {MAX_OCR_PAGES_PER_BATCH} PDF pages per batch.")
    try:
        pdf = PdfReader(io.BytesIO(data))
        if pdf.is_encrypted:
            raise SourceError("encrypted_pdf", "Upload an unlocked PDF.")
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise SourceError("pdf_page_limit", f"Limit PDFs to {MAX_PDF_PAGES} pages.")
        if start_page < 1 or end_page < start_page or end_page > len(pdf.pages):
            raise SourceError("pdf_batch_range", "The requested PDF page batch is invalid.")
        renderer = None
        sections: list[ExtractedSection] = []
        page_results: list[dict] = []
        try:
            for page_number in range(start_page, end_page + 1):
                page = pdf.pages[page_number - 1]
                try:
                    text = page.extract_text(extraction_mode="layout") or ""
                except TypeError:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""
                text = _normalise_pdf_text(text).strip()
                method = "text" if text else None
                warning = _page_text_warning(text, "text") if text else None
                status = "text-extracted" if text and not warning else "needs-review" if text else ""
                if not text:
                    if renderer is None:
                        import pypdfium2 as pdfium
                        renderer = pdfium.PdfDocument(io.BytesIO(data))
                    rendered_page = renderer[page_number - 1]
                    page_width, page_height = rendered_page.get_size()
                    scale = min(1.5, 2200 / max(page_width, page_height, 1))
                    image_bytes = io.BytesIO()
                    rendered_page.render(scale=scale).to_pil().save(image_bytes, format="PNG")
                    if _blank_page_image(image_bytes.getvalue()):
                        status = "blank"
                        page_results.append({"page_number": page_number, "status": status, "method": None, "warning": None, "text_chars": 0, "section_count": 0, "chapter_path": chapter_paths[page_number - 1] if page_number <= len(chapter_paths) else []})
                        continue
                    method = "vision"
                    if vision is None:
                        text = ""
                        warning = "vision_unavailable"
                    else:
                        try:
                            text = _normalise_pdf_text(vision(image_bytes.getvalue(), "image/png")).strip()
                            warning = _page_text_warning(text, "vision") if text else "empty_ocr"
                        except Exception:
                            text = ""
                            warning = "vision_unavailable"
                    status = "ocr-extracted" if text and not warning else "needs-review"

                chapter_path = chapter_paths[page_number - 1] if page_number <= len(chapter_paths) else []
                page_sections = _prose_sections(
                    text, filename, "pdf",
                    {"page": page_number, "reading_method": method or "text"},
                    markdown=True, all_caps=True,
                ) if text else []
                normalized_sections = []
                for section in page_sections:
                    local_heading_path = list(section.locator.get("heading_path", []))
                    full_heading_path = chapter_path + local_heading_path
                    locator = {
                        **section.locator,
                        "page": page_number,
                        "page_start": page_number,
                        "page_end": page_number,
                        "reading_method": method or "text",
                        "chapter_path": chapter_path,
                        "heading_path": full_heading_path,
                    }
                    title = " / ".join(full_heading_path) if chapter_path else section.title
                    normalized_sections.append(ExtractedSection(title, section.text, "pdf", locator))
                sections.extend(normalized_sections)
                page_results.append({
                    "page_number": page_number,
                    "status": status,
                    "method": method,
                    "warning": warning,
                    "text_chars": len(text),
                    "section_count": len(normalized_sections),
                    "chapter_path": chapter_path,
                })
        finally:
            if renderer is not None:
                renderer.close()
        return sections, page_results
    except SourceError:
        raise
    except Exception as exc:
        raise SourceError("unreadable_pdf", "The PDF page batch could not be read.") from exc


def merge_pdf_bookmark_sections(sections: list[ExtractedSection]) -> list[ExtractedSection]:
    """Join consecutive page fragments from one PDF bookmark path into one title."""
    merged: list[ExtractedSection] = []
    for section in sections:
        path = list(section.locator.get("heading_path", []))
        page = int(section.locator.get("page", 0))
        if path and merged:
            previous = merged[-1]
            previous_path = list(previous.locator.get("heading_path", []))
            previous_end = int(previous.locator.get("page_end", previous.locator.get("page", 0)))
            if path == previous_path and page <= previous_end + 1:
                locator = {**previous.locator, "page_end": max(previous_end, page), "section_char_end": len(previous.text) + 2 + len(section.text)}
                merged[-1] = ExtractedSection(previous.title, previous.text + "\n\n" + section.text, "pdf", locator)
                continue
        merged.append(section)
    return merged


def _docx_sections(data: bytes, filename: str) -> list[ExtractedSection]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    # DOCX is a zip container. Bound expansion before python-docx reads XML.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES or sum(entry.file_size for entry in entries) > MAX_OFFICE_EXPANDED_BYTES:
            raise SourceError('document_expansion_limit', 'The Word document exceeds the safe extraction limit.')
    document = Document(io.BytesIO(data))
    blocks: list[tuple[str, int, int]] = []
    paragraph_number = 0
    for item in document.iter_inner_content():
        paragraph_number += 1
        if isinstance(item, Paragraph):
            value = item.text.strip()
            if not value:
                continue
            style = item.style
            style_name = (style.name or '') if style is not None else ''
            style_id = (style.style_id or '') if style is not None else ''
            match = re.search(r'heading\s*([1-9])', style_name, re.IGNORECASE) or re.search(r'heading\s*([1-9])', style_id, re.IGNORECASE)
            if match:
                blocks.append((value, int(match.group(1)), paragraph_number))
            elif style_name.lower() == 'title' or style_id.lower() == 'title':
                blocks.append((value, 1, paragraph_number))
            else:
                blocks.append((value, 0, paragraph_number))
        elif isinstance(item, Table):
            for row_index, row in enumerate(item.rows, 1):
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append((' | '.join(cells), 0, paragraph_number * 1000 + row_index))
        if len(blocks) > MAX_DOCX_BLOCKS:
            raise SourceError('document_section_limit', 'This Word document has too many readable blocks.')
    sections: list[ExtractedSection] = []
    heading_path: list[tuple[int, str]] = []
    body: list[str] = []
    line_numbers: list[int] = []
    def flush():
        if not body:
            return
        content = '\n'.join(body)
        path = [value for _, value in heading_path]
        heading = ' / '.join(path) if path else Path(filename).stem[:160] or 'Source'
        sections.append(ExtractedSection(heading, content, 'text', {
            'filename': filename, 'heading': heading, 'heading_path': path,
            'line_start': line_numbers[0], 'line_end': line_numbers[-1],
            'docx_block_start': line_numbers[0], 'docx_block_end': line_numbers[-1],
        }))
        body.clear(); line_numbers.clear()
    for value, heading_level, source_position in blocks:
        if heading_level:
            flush()
            while heading_path and heading_path[-1][0] >= heading_level:
                heading_path.pop()
            heading_path.append((heading_level, value[:160]))
        else:
            body.append(value); line_numbers.append(source_position)
    flush()
    return sections


def _structured_sections(data: bytes, filename: str) -> list[ExtractedSection]:
    """Flatten JSON/YAML leaves while preserving key paths and source lines."""
    import yaml
    from yaml.nodes import MappingNode, ScalarNode, SequenceNode

    text = data.decode('utf-8-sig')
    try:
        root = yaml.compose(text, Loader=yaml.SafeLoader)
    except Exception as exc:
        raise SourceError('unreadable_structured_source', 'The JSON/YAML document could not be parsed safely.') from exc
    if root is None:
        return []
    sections: list[ExtractedSection] = []
    active_nodes: set[int] = set()
    def visit(node, path: str, depth: int = 0):
        if depth > MAX_STRUCTURED_DEPTH or len(sections) > MAX_STRUCTURED_VALUES:
            raise SourceError('structured_source_limit', 'The structured source is too deeply nested or has too many values.')
        identity = id(node)
        if identity in active_nodes:
            raise SourceError('structured_source_cycle', 'Recursive YAML aliases are not supported.')
        active_nodes.add(identity)
        try:
            if isinstance(node, MappingNode):
                for key_node, value_node in node.value:
                    key = str(getattr(key_node, 'value', '')).strip()
                    child = f'{path}.{key}' if path else key
                    visit(value_node, child, depth + 1)
            elif isinstance(node, SequenceNode):
                for index, value_node in enumerate(node.value):
                    visit(value_node, f'{path}[{index}]', depth + 1)
            elif isinstance(node, ScalarNode):
                raw = text[node.start_mark.index:node.end_mark.index]
                if not raw.strip():
                    return
                start, end = node.start_mark.index, node.end_mark.index
                sections.append(ExtractedSection(path or Path(filename).stem, raw, 'text', {
                    'filename': filename, 'heading': path or Path(filename).stem,
                    'heading_path': [path] if path else [], 'key_path': path,
                    'line_start': node.start_mark.line + 1, 'line_end': node.end_mark.line + 1,
                    'section_char_start': start, 'section_char_end': end,
                }))
        finally:
            active_nodes.remove(identity)
    visit(root, '')
    if len(sections) > MAX_STRUCTURED_VALUES:
        raise SourceError('section_limit', 'This structured source has too many values.')
    return sections


def _delimited_sections(data: bytes, filename: str, delimiter: str) -> list[ExtractedSection]:
    try:
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'), newline=''), delimiter=delimiter))
    except (UnicodeDecodeError, csv.Error) as exc:
        raise SourceError('unreadable_delimited_file', 'This CSV/TSV could not be read as UTF-8 rows.') from exc
    if sum(len(row) for row in rows) > MAX_TEXT_CELLS:
        raise SourceError('spreadsheet_limit', 'Limit parsed delimited files to 10,000 cells.')
    if not rows:
        return []
    headers = rows[0]
    columns = [header.strip() or f'Column {index + 1}' for index, header in enumerate(headers)]
    sections = []
    for row_number, row in enumerate(rows[1:], 2):
        pairs = [
            f'{columns[index] if index < len(columns) else f"Column {index + 1}"}: {value.strip()}'
            for index, value in enumerate(row) if value and value.strip()
        ]
        if not pairs:
            continue
        text = ' | '.join(pairs)
        sections.append(ExtractedSection(
            row[0][:160] if row and row[0].strip() else f'Row {row_number}', text, 'csv',
            {'filename': filename, 'row_start': row_number, 'row_end': row_number,
             'column_start': 1, 'column_end': max(1, len(row)),
             'columns': columns, 'headers': headers,
             'format': 'tsv' if delimiter == '\t' else 'csv',
             'row': True},
        ))
    return sections


def _xlsx_sections(data: bytes, filename: str) -> list[ExtractedSection]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES or sum(entry.file_size for entry in entries) > MAX_OFFICE_EXPANDED_BYTES:
                raise SourceError('spreadsheet_limit', 'The expanded workbook exceeds the safe extraction limit.')
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=filename.lower().endswith('.xlsm'))
        sections = []
        parsed_cells = 0
        for sheet in book:
            if sheet.max_row * sheet.max_column > MAX_TEXT_CELLS:
                raise SourceError('spreadsheet_limit', 'Limit parsed spreadsheets to 10,000 cells per sheet.')
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            labels = [str(value).strip() if value is not None else '' for value in headers]
            for row_number, row in enumerate(rows, 2):
                parsed_cells += len(row)
                if parsed_cells > MAX_TEXT_CELLS:
                    raise SourceError('spreadsheet_limit', 'Limit workbooks to 10,000 parsed cells.')
                pairs = [
                    f'{labels[index] if index < len(labels) and labels[index] else f"Column {index + 1}"}: {value}'
                    for index, value in enumerate(row) if value is not None and str(value).strip()
                ]
                if not pairs:
                    continue
                last_column = get_column_letter(max(1, len(row)))
                sections.append(ExtractedSection(
                    f'{sheet.title} - {row[0] if row and row[0] is not None else row_number}'[:160],
                    ' | '.join(pairs), 'spreadsheet',
                    {'filename': filename, 'sheet': sheet.title,
                     'cell_range': f'A{row_number}:{last_column}{row_number}',
                     'row_start': row_number, 'row_end': row_number,
                     'headers': labels, 'row': True},
                ))
        book.close()
        return sections
    except SourceError:
        raise
    except Exception as exc:
        raise SourceError('unreadable_spreadsheet', 'The spreadsheet could not be read.') from exc


def _pdf_sections(data: bytes, filename: str, vision=None) -> list[ExtractedSection]:
    from pypdf import PdfReader
    try:
        pdf = PdfReader(io.BytesIO(data))
        if pdf.is_encrypted:
            raise SourceError('encrypted_pdf', 'Upload an unlocked PDF.')
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise SourceError('pdf_page_limit', f'Limit PDFs to {MAX_PDF_PAGES} pages.')
        pages = []
        scanned = []
        for page_number, page in enumerate(pdf.pages, 1):
            try:
                text = page.extract_text(extraction_mode='layout') or ''
            except TypeError:
                try:
                    text = page.extract_text() or ''
                except Exception:
                    text = ''
            except Exception:
                text = ''
            text = _normalise_pdf_text(text).strip()
            if text:
                pages.append((page_number, text, 'text'))
            else:
                scanned.append(page_number)
        if scanned and (vision is None or len(scanned) > MAX_OCR_PAGES_PER_BATCH):
            code = 'ocr_required' if vision is None else 'pdf_ocr_page_limit'
            message = ('No text could be extracted from scanned PDF pages because the verified vision path is unavailable.' if vision is None
                       else f'This OCR batch has more than {MAX_OCR_PAGES_PER_BATCH} scanned pages.')
            raise SourceError(code, message)
        if scanned:
            import pypdfium2 as pdfium
            rendered = pdfium.PdfDocument(io.BytesIO(data))
            try:
                for page_number in scanned:
                    image = io.BytesIO()
                    rendered[page_number - 1].render(scale=2.0).to_pil().save(image, format='PNG')
                    try:
                        text = _normalise_pdf_text(vision(image.getvalue(), 'image/png')).strip()
                    except Exception as exc:
                        raise SourceError('vision_unavailable', 'The configured vision provider could not read a scanned PDF page.') from exc
                    if not text:
                        raise SourceError('unreadable_pdf_page', f'Page {page_number} needs a clearer scan.')
                    pages.append((page_number, text, 'vision'))
            finally:
                rendered.close()
        sections = []
        for page_number, text, method in sorted(pages):
            sections.extend(_prose_sections(
                text, filename, 'pdf',
                {'page': page_number, 'reading_method': method},
                markdown=True, all_caps=True,
            ))
            if len(sections) > 10000:
                raise SourceError('section_limit', 'This PDF has too many section headings.')
        return sections
    except SourceError:
        raise
    except Exception as exc:
        raise SourceError('unreadable_pdf', 'The PDF could not be read.') from exc


def extract_file(data: bytes, filename: str, mime: str = '', vision=None) -> SourceSnapshot:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise SourceError('unsupported_type', 'Use PDF, Markdown, text, Word, RST, CSV/TSV, XLSX, JSON/YAML, log, JPG or PNG.')
    if not data:
        raise SourceError('empty_source', 'This file is empty.')
    if suffix == '.pdf':
        sections = _pdf_sections(data, filename, vision)
    elif suffix in {'.md', '.markdown', '.txt', '.rst', '.log'}:
        if len(data) > MAX_SMALL_SOURCE_BYTES:
            raise SourceError('text_size_limit', 'Limit text sources to 2 MiB.')
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise SourceError('text_encoding', 'Use UTF-8 encoded text.') from exc
        sections = _prose_sections(
            text, filename, 'text', {},
            markdown=suffix in {'.md', '.markdown'}, rst=suffix == '.rst',
            all_caps=suffix == '.txt',
        )
    elif suffix == '.docx':
        sections = _docx_sections(data, filename)
    elif suffix in {'.csv', '.tsv'}:
        if len(data) > MAX_SMALL_SOURCE_BYTES:
            raise SourceError('text_size_limit', 'Limit CSV/TSV sources to 2 MiB.')
        sections = _delimited_sections(data, filename, '\t' if suffix == '.tsv' else ',')
    elif suffix in {'.xlsx', '.xlsm'}:
        sections = _xlsx_sections(data, filename)
    elif suffix in {'.json', '.yml', '.yaml'}:
        if len(data) > MAX_SMALL_SOURCE_BYTES:
            raise SourceError('text_size_limit', 'Limit structured sources to 2 MiB.')
        sections = _structured_sections(data, filename)
    else:
        if (suffix == '.png' and not data.startswith(b'\x89PNG\r\n\x1a\n')) or (suffix in {'.jpg', '.jpeg'} and not data.startswith(b'\xff\xd8\xff')):
            raise SourceError('invalid_photo', 'The photo does not match its file type.')
        try:
            from PIL import Image
            image=Image.open(io.BytesIO(data))
            width,height=image.size
            if width*height>MAX_IMAGE_PIXELS:
                raise SourceError('image_pixel_limit',f'Limit photos to {MAX_IMAGE_PIXELS//1_000_000} megapixels.')
            image.verify()
        except SourceError:
            raise
        except Exception as exc:
            raise SourceError('invalid_photo','The photo could not be decoded safely.') from exc
        if vision is None:
            raise SourceError('vision_unavailable', 'Photo reading needs a configured, verified vision provider. Text uploads remain available.')
        try:
            text = vision(data, 'image/png' if suffix == '.png' else 'image/jpeg')
        except Exception as exc:
            raise SourceError('vision_unavailable', 'The configured vision provider could not read this photo.') from exc
        if not text or not text.strip():
            raise SourceError('unreadable_photo', 'Retake this photo with clear lighting and readable text.')
        sections = [ExtractedSection(Path(filename).stem, text.strip(), 'photo', {'filename': filename, 'section': 1, 'reading_method': 'vision'})]

    if suffix in {'.pdf', '.md', '.markdown', '.txt', '.rst', '.docx', '.log'}:
        sections = [part for section in sections for part in coherent_sections(section)]
    if not sections:
        raise SourceError('empty_source', 'No readable study content was found.')
    if len(sections) > MAX_TEACHING_WINDOWS:
        raise SourceError('section_limit', 'This source has too many sections.')
    return SourceSnapshot(sections)
